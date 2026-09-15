"""Text extraction for document import (PDF / DOCX / PPTX / MD / TXT).

Used by the "create a course from a document" flow. Pure-Python parsers,
no system dependencies. Each parser is imported lazily so the API still
starts if a given library isn't installed.
"""

from __future__ import annotations

import io


class ExtractError(Exception):
    """Raised when a document can't be read or its type is unsupported."""


def _extract_pdf(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError:
        raise ExtractError("PDF support needs `pypdf`: pip install pypdf")
    try:
        reader = PdfReader(io.BytesIO(data))
        return "\n\n".join((page.extract_text() or "") for page in reader.pages)
    except Exception as e:  # noqa: BLE001
        raise ExtractError(f"Couldn't read the PDF: {e}")


def _extract_docx(data: bytes) -> str:
    try:
        import docx  # python-docx
    except ImportError:
        raise ExtractError("DOCX support needs `python-docx`: pip install python-docx")
    try:
        doc = docx.Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts)
    except Exception as e:  # noqa: BLE001
        raise ExtractError(f"Couldn't read the DOCX: {e}")


def _extract_pptx(data: bytes) -> str:
    try:
        from pptx import Presentation  # python-pptx
    except ImportError:
        raise ExtractError("PPTX support needs `python-pptx`: pip install python-pptx")
    try:
        prs = Presentation(io.BytesIO(data))
        parts: list[str] = []
        for i, slide in enumerate(prs.slides, 1):
            texts = [
                shape.text.strip()
                for shape in slide.shapes
                if getattr(shape, "has_text_frame", False) and shape.text.strip()
            ]
            if texts:
                parts.append(f"# Slide {i}\n" + "\n".join(texts))
        return "\n\n".join(parts)
    except Exception as e:  # noqa: BLE001
        raise ExtractError(f"Couldn't read the PPTX: {e}")


def extract_text(filename: str, data: bytes) -> str:
    """Extract plain text from an uploaded document, dispatching on the
    file extension. Returns the text (stripped); raises ExtractError on an
    unsupported type or a parse failure."""
    name = (filename or "").lower()
    if name.endswith(".pdf"):
        text = _extract_pdf(data)
    elif name.endswith(".docx"):
        text = _extract_docx(data)
    elif name.endswith(".pptx"):
        text = _extract_pptx(data)
    elif name.endswith((".md", ".markdown", ".txt")):
        text = data.decode("utf-8", errors="replace")
    else:
        raise ExtractError(
            f"Unsupported file type: {filename}. Use PDF, DOCX, PPTX, MD, or TXT."
        )
    text = text.strip()
    if not text:
        raise ExtractError("No text could be extracted from that document.")
    return text
